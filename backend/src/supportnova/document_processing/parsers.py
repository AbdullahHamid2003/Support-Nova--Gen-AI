"""Document parsers (SRS Step 5): PDF and DOCX (mandatory), TXT / Markdown / CSV (optional).

Each parser produces a :class:`ParsedDocument` with ordered sections that preserve the
section number, heading, heading level and (for PDFs) the page span. Section detection
is format-aware - font size / weight in PDFs, heading styles in DOCX, "#" markers in
Markdown - and falls back to numbered-heading detection ("5.2 Delay Compensation")
so that previously unseen documents from the hidden evaluation pack are handled too.
"""

from __future__ import annotations

import csv
import io
import math
import re
from collections import Counter
from dataclasses import dataclass, field
from typing import Any

import yaml

from supportnova.core.errors import DocumentProcessingError

NUMBERED_HEADING = re.compile(r"^\s*(\d{1,2}(?:\.\d{1,2}){0,3})\.?\s+([A-Z0-9][^\n]{1,150})$")
PAGE_FOOTER = re.compile(r"^\s*(page\s+\d+(\s+of\s+\d+)?|\d+\s*/\s*\d+|-\s*\d+\s*-)\s*$", re.IGNORECASE)
HEADER_FOOTER_BAND = 0.1  # share of the page height, top and bottom, where running headers and footers sit
META_PATTERNS: dict[str, re.Pattern[str]] = {
    "doc_id": re.compile(r"document\s+id\s*[:\-|]\s*([A-Z]{2,5}-[A-Z]{2,5}-\d{2,3})", re.IGNORECASE),
    "version": re.compile(r"\bversion\s*[:\-|]\s*v?(\d+(?:\.\d+){0,2})", re.IGNORECASE),
    "status": re.compile(r"\bstatus\s*[:\-|]\s*(active|previous|superseded|draft)\b", re.IGNORECASE),
    "effective_date": re.compile(r"effective(?:\s+date)?\s*[:\-|]\s*(\d{4}-\d{2}-\d{2})", re.IGNORECASE),
    "expiry_date": re.compile(r"expiry(?:\s+date)?\s*[:\-|]\s*(\d{4}-\d{2}-\d{2}|none|n/a)", re.IGNORECASE),
    "owner": re.compile(r"\bowner\s*[:\-]\s*([A-Za-z&,' ]{3,60}?)(?:\s*\||\s*$)", re.IGNORECASE | re.MULTILINE),
}


@dataclass
class ParsedSection:
    section_id: str
    heading: str
    level: int
    text: str = ""
    page_start: int | None = None
    page_end: int | None = None


@dataclass
class ParsedDocument:
    title: str
    file_format: str
    page_count: int | None
    sections: list[ParsedSection]
    full_text: str
    detected_metadata: dict[str, Any] = field(default_factory=dict)
    front_matter: dict[str, Any] = field(default_factory=dict)


def detect_metadata(text: str) -> dict[str, Any]:
    head = text[:3000]
    found: dict[str, Any] = {}
    for key, pat in META_PATTERNS.items():
        m = pat.search(head)
        if m:
            value = m.group(1).strip()
            if key == "status":
                value = value.capitalize()
            if key == "expiry_date" and value.lower() in ("none", "n/a"):
                value = None
            found[key] = value
    return found


def _level(section_id: str) -> int:
    return section_id.count(".") + 1


def _finalize(sections: list[ParsedSection]) -> list[ParsedSection]:
    for s in sections:
        s.text = re.sub(r"[ \t]+", " ", s.text).strip()
        s.text = re.sub(r"\n{3,}", "\n\n", s.text)
    return sections


def _sections_from_lines(lines: list[tuple[str, int | None, bool]], *,
                         require_hint: bool = True) -> tuple[list[ParsedSection], str]:
    """Build sections from (text, page, is_heading_hint) lines.

    A line is a section heading when it looks like "5.2 Delay Compensation" AND the format
    marks it as a heading (bold/larger font in PDF, heading style in DOCX, "#" in Markdown).
    Plain-text documents without markup use the numbering pattern alone (require_hint=False).
    """
    sections: list[ParsedSection] = []
    preamble: list[str] = []
    current: ParsedSection | None = None
    for text, page, heading_hint in lines:
        m = NUMBERED_HEADING.match(text)
        if m and (heading_hint or not require_hint) and not text.rstrip().endswith((".", ",", ";", ":")):
            current = ParsedSection(m.group(1), m.group(2).strip(), _level(m.group(1)), "", page, page)
            sections.append(current)
            continue
        if current is None:
            preamble.append(text)
        else:
            current.text += text + "\n"
            if page is not None:
                current.page_end = page
    return sections, "\n".join(preamble)


def _unnumbered_sections(lines: list[tuple[str, int | None, bool]]) -> list[ParsedSection]:
    """Fallback for documents without numbered headings: split at heading-styled lines."""
    sections: list[ParsedSection] = []
    current: ParsedSection | None = None
    for text, page, heading_hint in lines:
        if heading_hint and len(text) <= 100:
            current = ParsedSection(str(len(sections) + 1), text.strip(), 1, "", page, page)
            sections.append(current)
        elif current is None:
            current = ParsedSection("1", "Introduction", 1, text + "\n", page, page)
            sections.append(current)
        else:
            current.text += text + "\n"
            if page is not None:
                current.page_end = page
    return sections


# --------------------------------------------------------------------------- PDF
def _page_tables(page: Any) -> list[tuple[tuple[float, float, float, float], list[str]]]:
    """Tables on a PDF page as (bbox, rows); each row becomes "cell | cell | cell" so the columns stay together."""
    try:
        found = page.find_tables().tables
    except Exception:  # table detection is best effort - the page is then read as plain text lines
        return []
    tables = []
    for table in found:
        rows = []
        for row in table.extract():
            cells = [" ".join((cell or "").split()) or "-" for cell in row]
            if any(cell != "-" for cell in cells):
                rows.append(" | ".join(cells))
        if rows:
            tables.append((tuple(table.bbox), rows))
    return tables


def parse_pdf(data: bytes) -> ParsedDocument:
    try:
        import pymupdf as fitz
    except ImportError as exc:  # pragma: no cover
        raise DocumentProcessingError("PDF support requires PyMuPDF.") from exc
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # corrupted / encrypted
        raise DocumentProcessingError("The PDF could not be opened (corrupted or unsupported).") from exc
    if doc.needs_pass:
        raise DocumentProcessingError("Password-protected PDFs are not supported.")
    raw_lines: list[tuple[str, int, float, bool, bool]] = []  # text, page, font size, bold, in the header/footer band
    sizes: Counter[float] = Counter()
    for page_no, page in enumerate(iter(doc), start=1):  # Document iterates via __getitem__; iter() exposes Page to mypy
        height = page.rect.height or 1.0
        tables = _page_tables(page)
        emitted: set[int] = set()
        blocks = page.get_text("dict").get("blocks", [])
        for block in blocks:
            for line in block.get("lines", []):
                spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
                if not spans:
                    continue
                x0, y0, x1, y1 = line["bbox"]
                cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
                hit = next((i for i, (b, _rows) in enumerate(tables) if b[0] <= cx <= b[2] and b[1] <= cy <= b[3]), None)
                if hit is not None:  # a table cell: the whole table is emitted once, one line per row
                    if hit not in emitted:
                        emitted.add(hit)
                        raw_lines.extend((row, page_no, 0.0, False, False) for row in tables[hit][1])
                    continue
                text = " ".join(s["text"].strip() for s in spans).strip()
                size = round(max(s.get("size", 0) for s in spans), 1)
                bold = all((s.get("flags", 0) & 16) or "bold" in s.get("font", "").lower() for s in spans)
                band = y0 < height * HEADER_FOOTER_BAND or y1 > height * (1 - HEADER_FOOTER_BAND)
                raw_lines.append((text, page_no, size, bold, band))
                sizes[size] += len(text)
    page_count = doc.page_count
    doc.close()
    if not raw_lines:
        raise DocumentProcessingError("The PDF contains no extractable text (scanned images are not supported).")
    body_size = sizes.most_common(1)[0][0] if sizes else 10.0
    # running headers/footers: the same line in the top or bottom band of most pages. Body text that repeats
    # (a department name in a table, a recurring phrase) is content and is kept.
    line_pages: dict[str, set[int]] = {}
    for text, page_no, _s, _b, band in raw_lines:
        if band:
            line_pages.setdefault(text, set()).add(page_no)
    min_pages = max(2, math.ceil(page_count * 0.6))
    repeated = {t for t, pages in line_pages.items() if page_count > 1 and len(pages) >= min_pages}
    lines: list[tuple[str, int | None, bool]] = []
    title = ""
    for text, page_no, size, bold, band in raw_lines:
        if (band and text in repeated) or PAGE_FOOTER.match(text):
            continue
        if not title and page_no == 1 and size >= body_size + 4:
            title = text
            continue
        heading_hint = bold or size >= body_size + 1.5
        lines.append((text, page_no, heading_hint))
    # merge wrapped heading continuation lines is unnecessary for our renderer; join paragraphs
    sections, preamble = _sections_from_lines(lines)
    if not sections:
        sections = _unnumbered_sections(lines)
    full_text = "\n".join(t for t, _p, _h in lines)
    return ParsedDocument(title=title or (lines[0][0] if lines else "Untitled"), file_format="pdf", page_count=page_count,
                          sections=_finalize(sections), full_text=full_text,
                          detected_metadata=detect_metadata(preamble + "\n" + full_text))


# -------------------------------------------------------------------------- DOCX
def parse_docx(data: bytes) -> ParsedDocument:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentProcessingError("DOCX support requires python-docx.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentProcessingError("The DOCX file could not be opened (corrupted or unsupported).") from exc
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    lines: list[tuple[str, int | None, bool]] = []
    title = document.core_properties.title or ""
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style.lower() == "title":
                title = title or text
                continue
            heading_hint = style.lower().startswith("heading") or (bool(para.runs) and all(r.bold for r in para.runs if r.text.strip()))
            if style.lower().startswith("list"):
                text = "- " + text
            lines.append((text, None, heading_hint))
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                deduped = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]
                lines.append((" | ".join(deduped), None, False))
    if not lines:
        raise DocumentProcessingError("The DOCX document contains no text.")
    sections, preamble = _sections_from_lines(lines)
    if not sections:
        sections = _unnumbered_sections(lines)
    full_text = "\n".join(t for t, _p, _h in lines)
    meta = detect_metadata(preamble + "\n" + full_text)
    comments = document.core_properties.comments or ""
    for part in comments.split(";"):
        if "=" in part:
            k, v = part.split("=", 1)
            key = {"effective": "effective_date", "expiry": "expiry_date"}.get(k.strip(), k.strip())
            meta.setdefault(key, v.strip())
    if document.core_properties.subject and re.match(r"^[A-Z]{2,5}-[A-Z]{2,5}-\d{2,3}$", document.core_properties.subject):
        meta.setdefault("doc_id", document.core_properties.subject)
    return ParsedDocument(title=title or lines[0][0], file_format="docx", page_count=None, sections=_finalize(sections),
                          full_text=full_text, detected_metadata=meta)


# ---------------------------------------------------------------------- MD / TXT
_FRONT = re.compile(r"^---\s*\n(.*?)\n---\s*\n", re.DOTALL)
_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")


def parse_text(data: bytes, file_format: str) -> ParsedDocument:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    front: dict[str, Any] = {}
    m = _FRONT.match(text)
    if m:
        try:
            front = yaml.safe_load(m.group(1)) or {}
        except yaml.YAMLError:
            front = {}
        text = text[m.end():]
    lines: list[tuple[str, int | None, bool]] = []
    title = str(front.get("title") or "")
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        hm = _MD_HEADING.match(line)
        if hm:
            heading = hm.group(2).strip()
            if len(hm.group(1)) == 1 and not title:
                title = heading
                continue
            if len(hm.group(1)) == 1:
                continue
            lines.append((heading, None, True))
        else:
            lines.append((line.lstrip("*").rstrip("*") if line.startswith("**") else line, None, False))
    if not lines:
        raise DocumentProcessingError("The document contains no text.")
    has_markup = any(hint for _t, _p, hint in lines)
    sections, preamble = _sections_from_lines(lines, require_hint=has_markup)
    if not sections:
        sections = _unnumbered_sections(lines)
    full_text = "\n".join(t for t, _p, _h in lines)
    meta = detect_metadata(preamble + "\n" + full_text)
    for key in ("doc_id", "title", "doc_type", "version", "status", "effective_date", "expiry_date", "owner_department",
                "topics", "supersedes"):
        if front.get(key) not in (None, ""):
            meta[key] = str(front[key]) if key not in ("topics",) else front[key]
    return ParsedDocument(title=title or lines[0][0], file_format=file_format, page_count=None,
                          sections=_finalize(sections), full_text=full_text, detected_metadata=meta, front_matter=front)


def parse_csv(data: bytes) -> ParsedDocument:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [r for r in reader if any(c.strip() for c in r)]
    if not rows:
        raise DocumentProcessingError("The CSV file contains no rows.")
    header, body = rows[0], rows[1:] or rows
    sections = []
    for i, row in enumerate(body, start=1):
        heading = row[0].strip()[:150] or f"Row {i}"
        rest = "; ".join(f"{header[j] if j < len(header) else 'col' + str(j)}: {v}" for j, v in enumerate(row[1:], start=1) if v)
        sections.append(ParsedSection(str(i), heading, 1, rest))
    return ParsedDocument(title=header[0] if header else "CSV document", file_format="csv", page_count=None,
                          sections=sections, full_text=text, detected_metadata={})


def parse_document(data: bytes, file_format: str) -> ParsedDocument:
    fmt = file_format.lower()
    if fmt == "pdf":
        return parse_pdf(data)
    if fmt == "docx":
        return parse_docx(data)
    if fmt in ("md", "txt"):
        return parse_text(data, fmt)
    if fmt == "csv":
        return parse_csv(data)
    raise DocumentProcessingError(f"Unsupported document format '{file_format}'.")
