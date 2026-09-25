"""Report rendering to CSV, Excel and PDF (SRS Step 66, Deliverables 8 and 10).

One neutral `Report` model (sections of paragraphs, metrics, bullets and tables) renders to every
format, so the CSV, XLSX and PDF of a report always carry the same numbers. Untrusted text is
defended at the output boundary: CSV/XLSX cells that would be interpreted as spreadsheet formulas
are neutralised, control characters are stripped, and PDF text is transliterated to the core-font
character set.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

from fpdf import FPDF, XPos, YPos
from fpdf.fonts import FontFace

from supportnova.core.timeutil import utcnow

MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "json": "application/json",
}
FORMATS = tuple(MEDIA_TYPES)
BRAND = "SupportNova - ResponseX Intelligence"
FICTION_NOTE = "Lumora Home Technologies is a fictional organisation; all customers, orders and policies are synthetic demo data."

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")
_PDF_MAP = str.maketrans({
    "—": "-", "–": "-", "‒": "-", "−": "-", "‘": "'", "’": "'", "‚": ",", "“": '"',
    "”": '"', "…": "...", "•": "*", "·": "-", "→": "->", "←": "<-", "≥": ">=", "≤": "<=",
    "≠": "!=", " ": " ", " ": " ", "​": "", "✓": "v", "✗": "x", "×": "x", "€": "EUR",
})


@dataclass
class Table:
    columns: list[str]
    rows: list[list[Any]]
    title: str = ""
    widths: list[float] | None = None


@dataclass
class Section:
    heading: str
    paragraphs: list[str] = field(default_factory=list)
    metrics: list[tuple[str, Any]] = field(default_factory=list)
    bullets: list[str] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)


@dataclass
class Report:
    title: str
    sections: list[Section]
    subtitle: str = ""
    meta: dict[str, Any] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)
    banner: str | None = None  # an important notice, rendered prominently

    @property
    def generated_at(self) -> str:
        return str(self.meta.get("generated_at") or utcnow().strftime("%Y-%m-%d %H:%M UTC"))


# ---------------------------------------------------------------------------------------------- values
def display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return f"{value:.4f}".rstrip("0").rstrip(".") if abs(value) < 1 else f"{value:,.2f}".rstrip("0").rstrip(".")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, (list, tuple, set)):
        return "; ".join(display(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, default=str, ensure_ascii=False)
    return _CONTROL.sub("", str(value))


def csv_safe(text: str) -> str:
    """Neutralise spreadsheet formula injection (OWASP CSV injection)."""
    return "'" + text if text.startswith(_FORMULA_PREFIX) else text


def pdf_text(value: Any, limit: int | None = None) -> str:
    text = display(value).translate(_PDF_MAP)
    if limit and len(text) > limit:
        text = text[: limit - 3].rstrip() + "..."
    return text.encode("latin-1", errors="replace").decode("latin-1")


# ---------------------------------------------------------------------------------------------- csv
def _render_csv(report: Report) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    tables = [t for s in report.sections for t in s.tables]
    simple = len(tables) == 1 and not any(s.metrics or s.paragraphs or s.bullets for s in report.sections)
    if simple:
        w.writerow(tables[0].columns)
        for row in tables[0].rows:
            w.writerow([csv_safe(display(v)) for v in row])
    else:
        w.writerow([report.title])
        if report.subtitle:
            w.writerow([csv_safe(report.subtitle)])
        w.writerow(["Generated", report.generated_at])
        for k, v in report.meta.items():
            if k != "generated_at":
                w.writerow([k, csv_safe(display(v))])
        for s in report.sections:
            w.writerow([])
            w.writerow([f"## {s.heading}"])
            for label, value in s.metrics:
                w.writerow([label, csv_safe(display(value))])
            for t in s.tables:
                if t.title:
                    w.writerow([f"# {t.title}"])
                w.writerow(t.columns)
                for row in t.rows:
                    w.writerow([csv_safe(display(v)) for v in row])
        for note in report.notes:
            w.writerow([csv_safe(note)])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


# ---------------------------------------------------------------------------------------------- xlsx
def _render_xlsx(report: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E3A5F")
    title_font = Font(bold=True, size=14)
    used: set[str] = set()

    def sheet_name(name: str) -> str:
        base = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:28] or "Sheet"
        candidate, n = base, 2
        while candidate.lower() in used:
            candidate, n = f"{base[:25]} {n}", n + 1
        used.add(candidate.lower())
        return candidate

    def put(ws: Any, row: int, col: int, value: Any) -> Any:
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            cell = ws.cell(row=row, column=col, value=value)
        else:
            text = ILLEGAL_CHARACTERS_RE.sub("", display(value))
            cell = ws.cell(row=row, column=col, value=text)
            if text.startswith("="):
                cell.data_type = "s"  # never let untrusted text become a formula
        return cell

    summary = wb.active
    summary.title = sheet_name("Summary")
    summary.cell(row=1, column=1, value=report.title).font = title_font
    r = 2
    if report.subtitle:
        put(summary, r, 1, report.subtitle)
        r += 1
    if report.banner:
        put(summary, r, 1, report.banner).font = Font(bold=True, color="9C2A00")
        r += 1
    put(summary, r, 1, "Generated")
    put(summary, r, 2, report.generated_at)
    r += 1
    for k, v in report.meta.items():
        if k != "generated_at":
            put(summary, r, 1, k)
            put(summary, r, 2, v)
            r += 1
    for s in report.sections:
        if not (s.metrics or s.paragraphs or s.bullets):
            continue
        r += 1
        summary.cell(row=r, column=1, value=s.heading).font = Font(bold=True, size=12)
        r += 1
        for label, value in s.metrics:
            put(summary, r, 1, label)
            put(summary, r, 2, value)
            r += 1
        for p in s.paragraphs + [f"- {b}" for b in s.bullets]:
            put(summary, r, 1, p).alignment = Alignment(wrap_text=False)
            r += 1
    for note in [*report.notes, FICTION_NOTE]:
        r += 1
        put(summary, r, 1, note).font = Font(italic=True, color="555555")
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 60

    for s in report.sections:
        for t in s.tables:
            ws = wb.create_sheet(sheet_name(t.title or s.heading))
            for c, name in enumerate(t.columns, start=1):
                cell = put(ws, 1, c, name)
                cell.font, cell.fill = head_font, head_fill
            for ri, row in enumerate(t.rows, start=2):
                for c, value in enumerate(row, start=1):
                    put(ws, ri, c, value)
            ws.freeze_panes = "A2"
            if t.rows:
                ws.auto_filter.ref = f"A1:{get_column_letter(len(t.columns))}{len(t.rows) + 1}"
            for c, name in enumerate(t.columns, start=1):
                sample = [len(display(row[c - 1])) for row in t.rows[:200] if c - 1 < len(row)]
                ws.column_dimensions[get_column_letter(c)].width = max(10, min(60, max([len(name), *sample]) + 2))
    if len(wb.sheetnames) > 1 and not any(s.metrics or s.paragraphs or s.bullets for s in report.sections):
        wb.move_sheet(summary, offset=len(wb.sheetnames) - 1)  # data first, "About" last for plain exports
        summary.title = sheet_name("About")
        wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------- pdf
class _ReportPDF(FPDF):
    def __init__(self, report: Report, orientation: str) -> None:
        super().__init__(orientation=orientation, unit="mm", format="A4")
        self.report = report
        self.set_auto_page_break(auto=True, margin=16)
        self.set_margins(14, 14, 14)
        self.set_title(pdf_text(report.title))
        self.set_author(BRAND)
        self.set_creator(BRAND)

    def header(self) -> None:
        self.set_font("Helvetica", "B", 8)
        self.set_text_color(30, 58, 95)
        self.cell(0, 5, pdf_text(BRAND), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_draw_color(30, 58, 95)
        self.set_line_width(0.4)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.ln(3)
        self.set_text_color(0, 0, 0)

    def footer(self) -> None:
        self.set_y(-11)
        self.set_font("Helvetica", "", 7)
        self.set_text_color(110, 110, 110)
        self.cell(0, 4, pdf_text(f"{self.report.title} - generated {self.report.generated_at} - fictional demo data"), align="L")
        self.set_x(self.l_margin)
        self.cell(0, 4, f"Page {self.page_no()}/{{nb}}", align="R")
        self.set_text_color(0, 0, 0)


def _widths(t: Table) -> list[float]:
    if t.widths:
        return t.widths
    out = []
    for i, name in enumerate(t.columns):
        lengths = sorted(len(display(r[i])) for r in t.rows[:300] if i < len(r)) or [0]
        typical = lengths[int(len(lengths) * 0.8)] if lengths else 0
        out.append(float(max(5, min(42, max(len(name) * 0.8, typical)))))
    return out


def _pdf_table(pdf: _ReportPDF, t: Table, *, font_size: float = 7.2, cell_limit: int = 420) -> None:
    if t.title:
        pdf.set_font("Helvetica", "B", 9)
        pdf.cell(0, 6, pdf_text(t.title), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    if not t.rows:
        pdf.set_font("Helvetica", "I", 8)
        pdf.cell(0, 5, "No records.", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(2)
        return
    pdf.set_font("Helvetica", "", font_size)
    headings = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=(30, 58, 95))
    with pdf.table(col_widths=tuple(_widths(t)), headings_style=headings, line_height=font_size * 0.52,
                   text_align="LEFT", borders_layout="HORIZONTAL_LINES", cell_fill_color=(244, 247, 251),
                   cell_fill_mode="ROWS", repeat_headings=1) as table:
        head = table.row()
        for name in t.columns:
            head.cell(pdf_text(name))
        for row in t.rows:
            r = table.row()
            for i in range(len(t.columns)):
                r.cell(pdf_text(row[i] if i < len(row) else "", cell_limit))
    pdf.ln(3)


def _render_pdf(report: Report) -> bytes:
    widest = max((len(t.columns) for s in report.sections for t in s.tables), default=0)
    pdf = _ReportPDF(report, "L" if widest > 7 else "P")
    pdf.alias_nb_pages()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 17)
    pdf.multi_cell(0, 8, pdf_text(report.title), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    if report.subtitle:
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(70, 70, 70)
        pdf.multi_cell(0, 5, pdf_text(report.subtitle), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_text_color(0, 0, 0)
    pdf.set_font("Helvetica", "", 8)
    meta = [f"Generated: {report.generated_at}"] + [f"{k}: {display(v)}" for k, v in report.meta.items() if k != "generated_at"]
    pdf.multi_cell(0, 4.2, pdf_text("   |   ".join(meta)), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    if report.banner:
        pdf.ln(2)
        pdf.set_fill_color(255, 243, 205)
        pdf.set_font("Helvetica", "B", 8.5)
        pdf.multi_cell(0, 5, pdf_text(report.banner), fill=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(3)
    for s in report.sections:
        if pdf.get_y() > pdf.h - 40:
            pdf.add_page()
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(30, 58, 95)
        pdf.cell(0, 7, pdf_text(s.heading), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_text_color(0, 0, 0)
        if s.metrics:
            _pdf_metrics(pdf, s.metrics)
        pdf.set_font("Helvetica", "", 8.5)
        for p in s.paragraphs:
            pdf.multi_cell(0, 4.4, pdf_text(p), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(1)
        for b in s.bullets:
            pdf.set_x(pdf.l_margin + 3)
            pdf.multi_cell(0, 4.4, pdf_text(f"- {b}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        if s.bullets:
            pdf.ln(1)
        for t in s.tables:
            _pdf_table(pdf, t)
        pdf.ln(1)
    pdf.set_font("Helvetica", "I", 7.5)
    pdf.set_text_color(90, 90, 90)
    for note in [*report.notes, FICTION_NOTE]:
        pdf.multi_cell(0, 4, pdf_text(note), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    return bytes(pdf.output())


def _pdf_metrics(pdf: _ReportPDF, metrics: list[tuple[str, Any]]) -> None:
    per_row = 4 if pdf.w > 250 else 3
    gap = 3.0
    width = (pdf.w - pdf.l_margin - pdf.r_margin - gap * (per_row - 1)) / per_row
    for i in range(0, len(metrics), per_row):
        if pdf.get_y() > pdf.h - 30:
            pdf.add_page()
        y = pdf.get_y()
        for j, (label, value) in enumerate(metrics[i:i + per_row]):
            x = pdf.l_margin + j * (width + gap)
            pdf.set_fill_color(244, 247, 251)
            pdf.rect(x, y, width, 13, style="F")
            pdf.set_xy(x + 2, y + 1.5)
            pdf.set_font("Helvetica", "", 7)
            pdf.set_text_color(90, 90, 90)
            pdf.cell(width - 4, 3.5, pdf_text(label, 60))
            pdf.set_xy(x + 2, y + 5.5)
            pdf.set_font("Helvetica", "B", 10.5)
            pdf.set_text_color(20, 20, 20)
            pdf.cell(width - 4, 6, pdf_text(value if value not in (None, "") else "-", 40))
        pdf.set_text_color(0, 0, 0)
        pdf.set_xy(pdf.l_margin, y + 15)
    pdf.ln(1)


# ---------------------------------------------------------------------------------------------- api
def render(report: Report, fmt: str) -> tuple[bytes, str, str]:
    fmt = (fmt or "csv").lower()
    if fmt == "csv":
        return _render_csv(report), MEDIA_TYPES["csv"], "csv"
    if fmt in ("xlsx", "excel"):
        return _render_xlsx(report), MEDIA_TYPES["xlsx"], "xlsx"
    if fmt == "pdf":
        return _render_pdf(report), MEDIA_TYPES["pdf"], "pdf"
    if fmt == "json":
        payload = {"title": report.title, "subtitle": report.subtitle, "meta": report.meta, "banner": report.banner,
                   "sections": [{"heading": s.heading, "paragraphs": s.paragraphs, "bullets": s.bullets,
                                 "metrics": [{"label": k, "value": v} for k, v in s.metrics],
                                 "tables": [{"title": t.title, "columns": t.columns, "rows": t.rows} for t in s.tables]}
                                for s in report.sections], "notes": report.notes}
        return json.dumps(payload, default=str, indent=2).encode("utf-8"), MEDIA_TYPES["json"], "json"
    from supportnova.core.errors import ValidationFailed
    raise ValidationFailed(f"Unsupported export format '{fmt}'. Use one of: csv, xlsx, pdf, json.")


def tabular(title: str, columns: list[str], rows: list[list[Any]], fmt: str, *, subtitle: str = "",
            meta: dict[str, Any] | None = None) -> tuple[bytes, str, str]:
    report = Report(title=title, subtitle=subtitle, meta=meta or {}, sections=[Section(heading=title, tables=[Table(columns, rows)])])
    return render(report, fmt)


# ---------------------------------------------------------------------------------------------- case report
# how the final classification was reached (validated_decision.classification.source)
_CLASSIFICATION_SOURCE = {"python_rules": "rules", "reviewer_override": "reviewer", "ai_unconfirmed": "AI, awaiting review",
                          "python_low_confidence": "rules, low confidence, awaiting review"}


def _field_label(field: Any) -> str:
    return str(field or "").replace("_", " ").capitalize()


def _reason_text(r: dict[str, Any]) -> str:
    """A review reason as "name: detail" (stored reasons carry name/detail; older ones may carry message)."""
    detail = r.get("detail") or r.get("message")
    return f"{r.get('name') or r.get('code')}" + (f": {detail}" if detail else "")


def _kv_table(title: str, pairs: list[tuple[str, Any]]) -> Table:
    return Table(["Field", "Value"], [[k, v] for k, v in pairs if v not in (None, "", [], {})], title=title, widths=[30, 110])


def complaint_case_pdf(detail: dict[str, Any], history: list[dict[str, Any]]) -> bytes:
    """Complete case report: complaint, validated final decision, AI-vs-rules comparison,
    validation checks, evidence, response, escalation, review and timeline."""
    analysis = detail.get("analysis") or {}
    validation = detail.get("validation") or {}
    final = validation.get("validated_decision") or {}
    cls = final.get("classification") or {}
    elig = final.get("eligibility") or {}
    esc = final.get("escalation") or {}
    sections: list[Section] = []
    sections.append(Section("Complaint", metrics=[
        ("Status", detail.get("status")), ("Verification", detail.get("verification_status")),
        ("Verification score", detail.get("verification_score")), ("Priority", detail.get("priority")),
        ("Urgency", detail.get("urgency")), ("Department", detail.get("department")),
        ("Escalation", detail.get("escalation_level")), ("SLA", detail.get("sla_state")),
    ], tables=[_kv_table("Case facts", [
        ("Customer", f"{detail.get('customer_name') or ''} ({detail.get('customer_ref') or 'n/a'}, {detail.get('customer_type')})"),
        ("Channel", detail.get("channel")), ("Submitted", detail.get("complaint_date")), ("Product", detail.get("product_text")),
        ("Order reference", detail.get("order_ref")), ("Transaction reference", detail.get("transaction_ref")),
        ("Previous complaint", detail.get("previous_complaint_ref")), ("Requested resolution", detail.get("requested_resolution")),
        ("Preferred contact", detail.get("preferred_contact")), ("Requested tone", detail.get("requested_tone")),
        ("Attachments", [a.get("file_name") for a in detail.get("attachments") or []]),
        ("Flags", [f for f, on in (("repeat", detail.get("is_repeat")), ("duplicate", detail.get("is_duplicate")),
                                   ("manipulation attempt detected", detail.get("injection_detected"))) if on]),
    ])], paragraphs=[f"Title: {detail.get('title')}", f"Description: {detail.get('description')}"]
        + ([f"Supporting information: {detail.get('supporting_info')}"] if detail.get("supporting_info") else [])))
    if final:
        sections.append(Section("Final decision (checked against the rules)", tables=[_kv_table("Decision", [
            ("Category / subcategory", f"{cls.get('category_name')} / {cls.get('subcategory_name')} "
                                       f"(source: {_CLASSIFICATION_SOURCE.get(cls.get('source') or '', cls.get('source'))})"),
            ("Secondary issues", cls.get("secondary")), ("Department", final.get("department")),
            ("Supporting departments", final.get("supporting_departments")),
            ("Urgency / impact / priority", f"{final.get('urgency')} / {final.get('impact')} / {final.get('priority')}"),
            ("Urgency sources", final.get("urgency_sources")),
            ("Escalation", f"{esc.get('level')} - " + "; ".join(f"{f.get('rule_id')} {f.get('reason')}" for f in esc.get("fired", [])[:5])),
            ("Refund / replacement / compensation", f"{elig.get('refund')} / {elig.get('replacement')} / {elig.get('compensation')}"),
            ("Compensation", f"{elig.get('compensation_type') or '-'} {elig.get('compensation_amount_usd') or ''}".strip()),
            ("Pending verification", elig.get("pending_rule_ids")), ("Selected rule", (final.get("selected_rule") or {}).get("rule_id")),
            ("Rule condition", (final.get("selected_rule") or {}).get("condition")),
            ("Required actions", [a if isinstance(a, str) else "any of " + "/".join(a.get("any_of", [])) for a in final.get("required_actions", [])]),
            ("Prohibited actions", final.get("prohibited_actions")), ("Policy references", final.get("policy_refs")),
            ("Missing information", [m.get("label") for m in final.get("missing_information", [])]),
            ("Clarification questions", final.get("clarification_questions")),
            ("Follow-up", (final.get("follow_up") or {}).get("type")),
            ("Timelines", [t.get("text") for t in final.get("timelines", [])]),
        ]), Table(["#", "Action", "Description", "Policy", "Source"],
                  [[i + 1, s.get("action_code"), s.get("description"), s.get("policy_ref"), s.get("source")]
                   for i, s in enumerate(final.get("resolution_steps", []))], title="Validated resolution steps",
                  widths=[6, 34, 70, 22, 12])],
            bullets=list((final.get("agent_guidance") or {}).get("validated", []))))
    comparison = (validation.get("comparison") or {}).get("rows", [])
    if comparison:
        sections.append(Section("AI vs rules", paragraphs=[
            f"Field agreement: {round(100 * (validation.get('comparison') or {}).get('agreement', 0))}%. "
            "The rules decide every rule-based field; mismatches are explained below."],
            tables=[Table(["Field", "AI", "Rules", "Result", "Basis"],
                          [[_field_label(r.get("field")), r.get("ai"), r.get("python"), r.get("match"), r.get("explanation")]
                           for r in comparison],
                          widths=[26, 38, 38, 14, 34])]))
    checks = validation.get("checks") or []
    if checks:
        sections.append(Section("Validation checks", metrics=[
            ("Decision", validation.get("decision")), ("Score", validation.get("score")),
            ("Failed", sum(c["status"] == "fail" for c in checks)), ("Warnings", sum(c["status"] == "warn" for c in checks)),
        ], tables=[Table(["Code", "Check", "Severity", "Result", "Detail"],
                         [[c["code"], c["name"], c["severity"], c["status"], c["message"]] for c in checks
                          if c["status"] != "not_applicable"], widths=[12, 40, 13, 12, 70])],
            bullets=[_reason_text(r) for r in validation.get("review_reasons") or []]))
    evidence = (analysis.get("retrieval") or {}).get("evidence") or []
    if evidence:
        sections.append(Section("Policy evidence", tables=[Table(
            ["Evidence", "Document", "Version", "Status", "Section", "Heading", "Score"],
            [[e.get("evidence_id"), e.get("doc_id"), e.get("version"), e.get("status"), e.get("section_id"), e.get("heading"),
              round(float(e.get("score") or 0), 3)] for e in evidence], widths=[14, 22, 12, 16, 12, 50, 12])]))
    responses = detail.get("responses") or []
    if responses:
        latest = responses[0]
        sections.append(Section("Customer response", metrics=[("Status", latest.get("status")), ("Tone", latest.get("tone")),
                                                              ("Source", latest.get("source")), ("Sent", latest.get("sent_at"))],
                                paragraphs=[f"Subject: {latest.get('subject')}", latest.get("body") or ""]))
    if detail.get("escalations"):
        sections.append(Section("Escalations", tables=[Table(["Level", "Source", "Rules", "Reason", "Status", "Created"],
            [[e.get("level"), e.get("source"), e.get("rule_ids"), e.get("reason"), e.get("status"), e.get("created_at")]
             for e in detail["escalations"]], widths=[30, 12, 22, 60, 12, 24])]))
    review = detail.get("review")
    if review:
        sections.append(Section("Manual review", metrics=[("Status", review.get("status")), ("Decision", review.get("final_decision"))],
                                tables=[Table(["Action", "Actor", "At", "Comment"],
                                              [[a.get("action"), a.get("actor"), a.get("at"), a.get("comment")] for a in review.get("actions", [])],
                                              widths=[18, 30, 30, 80])],
                                bullets=[_reason_text(r) for r in review.get("reasons") or []]))
    sections.append(Section("Complaint timeline", tables=[Table(["At", "Event", "Actor", "Detail"],
        [[h.get("at"), h.get("event_type"), h.get("actor"), h.get("message")] for h in history], widths=[30, 30, 30, 80])]))
    prompts = ", ".join(f"{v.get('key')} v{v.get('version')}" for v in (analysis.get("prompt_versions") or {}).values()
                        if isinstance(v, dict))
    trace = [f"AI provider: {analysis.get('provider')} / {analysis.get('model')}", f"Prompt versions: {prompts or '-'}",
             f"Rule set version: {analysis.get('ruleset_hash')}", f"Processing time: {analysis.get('total_latency_ms')} ms",
             f"AI attempts: {analysis.get('ai_attempts')}"] if analysis else ["No analysis has been run for this complaint."]
    sections.append(Section("Traceability", bullets=trace))
    report = Report(title=f"Case report {detail.get('complaint_ref')}", subtitle=str(detail.get("title") or ""), sections=sections,
                    meta={"Complaint": detail.get("complaint_ref"), "Category": detail.get("category"), "Subcategory": detail.get("subcategory")})
    return _render_pdf(report)
