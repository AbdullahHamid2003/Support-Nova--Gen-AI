"""Assemble the SupportNova Project Report from its chapter files and export it to Markdown, HTML, PDF and DOCX.

    backend/.venv/Scripts/python.exe documentation/_build/render_diagrams.py     (diagrams -> svg + png)
    backend/.venv/Scripts/python.exe documentation/_build/build_report.py        (md + html + pdf + docx)
    options: --check (validate only), --no-pdf, --no-docx

Inputs: documentation/chapters/NN-*.md (00 = executive summary), chapters/references.md, appendices/X-*.md,
the rendered diagrams and ../screenshots. Outputs in documentation/:
  SupportNova_Project_Report.md    GitHub-friendly Markdown (headings get GitHub anchors, TOC links work there)
  SupportNova_Project_Report.html  print layout used for the PDF (figures as SVG)
  SupportNova_Project_Report.pdf   A4, cover without header/footer, page-numbered TOC and lists, bookmarks
  SupportNova_Project_Report.docx  Word version (figures as PNG, TOC / figure / table lists as updatable fields)
  _build/report-check.txt          numbering, cross-reference and image audit
"""

from __future__ import annotations

import asyncio
import base64
import contextlib
import datetime as dt
import html
import re
import shutil
import subprocess
import sys
import tempfile
import time
from dataclasses import dataclass, field
from pathlib import Path

import httpx
import websockets

DOC = Path(__file__).resolve().parents[1]
ROOT = DOC.parent
sys.path.insert(0, str(DOC / "_build" / ".tools" / "pylib"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from render_diagrams import CDP, find_browser, free_port  # noqa: E402

VERSION = "1.0"
DATE = dt.date(2026, 9, 25)
TITLE = "SupportNova — Project Report"
PRINCIPLE = "GenAI proposes. Python validates. Ground truth decides."
META = [("Project", "SupportNova — Customer Complaint Resolution Intelligence"), ("Theme", "ResponseX Intelligence"),
        ("Category", "Generative AI PowerPlay"), ("Team / Project", "[Team name and members]"),
        ("Institution / Organization", "[Institution name]"), ("Version", VERSION), ("Date", f"{DATE:%d %B %Y}"),
        ("Document status", "Final — for submission")]
FIG_CAP = re.compile(r"^\*(Figure ([A-Z]|\d+)\.(\d+) — .+?)\*\s*$")
TAB_CAP = re.compile(r"^\*\*(Table ([A-Z]|\d+)\.(\d+) — .+?)\*\*\s*$")
IMG = re.compile(r"^!\[([^\]]*)\]\(([^)\s]+)\)\s*$")
HEAD = re.compile(r"^(#{1,6}) (.+?)\s*$")
FENCE = re.compile(r"^\s*(```|~~~)")


# ----------------------------------------------------------------------------------------------- model
@dataclass
class Heading:
    level: int
    text: str
    anchor: str


@dataclass
class Caption:
    kind: str       # Figure | Table
    label: str      # 4.1, B.2
    text: str       # full caption, "Figure 4.1 — ..."
    anchor: str
    part: str


@dataclass
class Report:
    parts: list[tuple[Path, str]] = field(default_factory=list)
    headings: list[Heading] = field(default_factory=list)
    figures: list[Caption] = field(default_factory=list)
    tables: list[Caption] = field(default_factory=list)
    issues: list[str] = field(default_factory=list)
    anchors: dict[tuple[int, int], str] = field(default_factory=dict)   # (part index, line index) -> heading anchor
    front_anchors: list[str] = field(default_factory=list)


_GH_STRIP = re.compile(r"[^\w\- ]", re.UNICODE)


class Slugger:
    """GitHub's heading-anchor algorithm (github-slugger), so the Markdown TOC links also work on GitHub."""

    def __init__(self) -> None:
        self.seen: dict[str, int] = {}

    def slug(self, text: str) -> str:
        t = re.sub(r"`([^`]*)`", r"\1", text)
        t = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", t).replace("*", "")
        base = _GH_STRIP.sub("", t.strip().lower()).replace(" ", "-")
        result = base
        while result in self.seen:
            self.seen[base] += 1
            result = f"{base}-{self.seen[base]}"
        self.seen[result] = 0
        return result


def git(*args: str) -> str:
    return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True).stdout.strip()


def part_files() -> list[Path]:
    chapters = sorted(DOC.glob("chapters/[0-9][0-9]-*.md"))
    refs = [DOC / "chapters" / "references.md"] if (DOC / "chapters" / "references.md").exists() else []
    return chapters + refs + sorted(DOC.glob("appendices/[A-Z]-*.md"))


FRONT = ["SupportNova — Customer Complaint Resolution Intelligence", "Document Control", "Table of Contents", "List of Figures",
         "List of Tables"]


# ----------------------------------------------------------------------------------------------- parse + audit
def load() -> Report:  # noqa: C901, PLR0912
    rep = Report()
    slugger = Slugger()
    rep.front_anchors = [slugger.slug(t) for t in FRONT]
    labels: dict[str, str] = {}
    for p in part_files():
        rep.parts.append((p, p.read_text(encoding="utf-8").replace("\r\n", "\n").strip() + "\n"))
    for pi, (p, text) in enumerate(rep.parts):
        lines = text.split("\n")
        m0 = re.match(r"^# (?:Chapter (\d+)|Appendix ([A-Z])) — ", lines[0])
        part_no = (m0.group(1) or m0.group(2)) if m0 else None
        if p.name != "references.md" and not m0:
            rep.issues.append(f"{p.name}: first line must be '# Chapter N — Title' or '# Appendix X — Title'")
        seq = {"Figure": 0, "Table": 0}
        in_code = False
        for li, ln in enumerate(lines):
            if FENCE.match(ln):
                in_code = not in_code
                continue
            if in_code:
                continue
            h = HEAD.match(ln)
            if h:
                anchor = slugger.slug(h.group(2))
                rep.anchors[(pi, li)] = anchor
                rep.headings.append(Heading(len(h.group(1)), h.group(2), anchor))
                if len(h.group(1)) == 1 and li > 0:
                    rep.issues.append(f"{p.name}: extra level-1 heading '{h.group(2)[:50]}' (use ## inside a chapter)")
                continue
            for kind, rx in (("Figure", FIG_CAP), ("Table", TAB_CAP)):
                cm = rx.match(ln.strip())
                if not cm:
                    continue
                label = f"{cm.group(2)}.{cm.group(3)}"
                anchor = f"{'fig' if kind == 'Figure' else 'tab'}-{cm.group(2).lower()}-{cm.group(3)}"
                key = f"{kind} {label}"
                if key in labels:
                    rep.issues.append(f"{p.name}: duplicate {key} (also in {labels[key]})")
                labels[key] = p.name
                if part_no and cm.group(2) != part_no:
                    rep.issues.append(f"{p.name}: {key} numbered outside its chapter/appendix {part_no}")
                seq[kind] += 1
                if int(cm.group(3)) != seq[kind]:
                    rep.issues.append(f"{p.name}: {key} out of sequence (expected {kind} {part_no}.{seq[kind]})")
                (rep.figures if kind == "Figure" else rep.tables).append(Caption(kind, label, cm.group(1), anchor, p.name))
                if kind == "Table":
                    nxt = next((x for x in lines[li + 1:li + 4] if x.strip()), "")
                    if not nxt.lstrip().startswith("|"):
                        rep.issues.append(f"{p.name}: {key} caption is not followed by a table")
            im = IMG.match(ln.strip())
            if im:
                src = im.group(2)
                target = (DOC / src).resolve()
                if not target.exists():
                    rep.issues.append(f"{p.name}: missing image {src}")
                if src.endswith(".svg") and not target.with_suffix(".png").exists():
                    rep.issues.append(f"{p.name}: missing PNG for {src} (needed by the DOCX)")
                nxt = next((x for x in lines[li + 1:li + 3] if x.strip()), "")
                cm = FIG_CAP.match(nxt.strip())
                if not cm:
                    rep.issues.append(f"{p.name}: image {src} has no figure caption line after it")
                elif im.group(1).strip() != cm.group(1).strip():
                    rep.issues.append(f"{p.name}: alt text differs from caption for {cm.group(1)[:50]}")
    whole = "\n".join(t for _p, t in rep.parts)
    known = {f"{c.kind} {c.label}" for c in rep.figures + rep.tables}
    for kind, num in sorted(set(re.findall(r"\b(Figure|Table) ((?:[A-Z]|\d+)\.\d+)\b", whole))):
        if f"{kind} {num}" not in known:
            rep.issues.append(f"cross-reference to missing {kind} {num}")
    chapters = {int(m) for m in re.findall(r"^# Chapter (\d+) — ", whole, re.M)}
    cited = {int(n) for n in re.findall(r"\bChapters? (\d+)\b", whole)}
    cited |= {int(n) for n in re.findall(r"\bChapters \d+(?:, \d+)*,? (?:and|to|or) (\d+)\b", whole)}
    for num in sorted(cited - chapters):
        rep.issues.append(f"cross-reference to missing Chapter {num}")
    appendices = set(re.findall(r"^# Appendix ([A-Z]) — ", whole, re.M))
    for letter in sorted(set(re.findall(r"\bAppendix ([A-Z])\b", whole)) - appendices):
        rep.issues.append(f"cross-reference to missing Appendix {letter}")
    for mmd in sorted((DOC / "diagrams").rglob("*.mmd")):
        if "_selftest" not in mmd.parts and mmd.with_suffix(".svg").relative_to(DOC).as_posix() not in whole:
            rep.issues.append(f"diagram not used in any chapter: {mmd.relative_to(DOC).as_posix()}")
    return rep


def doc_control(commit: str) -> list[str]:
    return ["| Field | Value |", "|---|---|",
            f"| Document Title | {TITLE} |",
            "| Project | SupportNova — Customer Complaint Resolution Intelligence (theme ResponseX Intelligence, category Generative AI PowerPlay) |",
            f"| Version | {VERSION} |", "| Status | Final — for submission |",
            "| Prepared By | [Team name / members] |", "| Reviewed By | [Reviewer name] |", f"| Date | {DATE:%d %B %Y} |",
            "| Source Specification | Aptech Limited, SupportNova — Customer Complaint Resolution Intelligence, Software Requirements "
            "Specification, Generative AI PowerPlay (issued to the teams; not redistributed in the repository) |",
            f"| Implementation Baseline | Repository `github.com/AbdullahHamid2003/Supoort-Nova--Gen-AI`, branch `main`, commit `{commit}` |",
            "| Purpose | Technical documentation of the implemented SupportNova system — requirements, architecture, pipelines, validation, "
            "security, testing and measured evidence — for competition submission and evaluation |"]


# ----------------------------------------------------------------------------------------------- markdown (public)
def public_markdown(rep: Report) -> str:
    out = [f"# {FRONT[0]}", "", "**Project Report**", "", "| Item | Value |", "|---|---|", *[f"| {k} | {v} |" for k, v in META], "",
           f"> **{PRINCIPLE}**", "", "*Lumora Home Technologies and all customers, orders, policies and complaints in this report are fictional.*",
           "", "# Document Control", "", *doc_control(git("rev-parse", "--short", "HEAD")), "", "# Table of Contents", ""]
    out += [("" if h.level == 1 else "  ") + f"- [{h.text}](#{h.anchor})" for h in rep.headings if h.level <= 2]
    out += ["", "# List of Figures", "", *[f"- {c.text}" for c in rep.figures], "", "# List of Tables", "", *[f"- {c.text}" for c in rep.tables], ""]
    return "\n".join(out) + "\n" + "\n".join(t for _p, t in rep.parts)


# ----------------------------------------------------------------------------------------------- html (print)
def inline_html(text: str) -> str:
    t = html.escape(text, quote=False)
    t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
    t = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", t)
    return re.sub(r"(?<![\w*])\*([^*\s][^*]*)\*(?![\w*])", r"<em>\1</em>", t)


def nav(entries: list[tuple[int, str, str]], cls: str) -> list[str]:
    return [f'<nav class="toc {cls}">', *[f'<a class="toc-{lvl}" href="#{a}"><span class="toc-t">{inline_html(t)}</span>'
                                          f'<span class="toc-d"></span><span class="toc-p" data-a="{a}">000</span></a>' for lvl, t, a in entries],
            "</nav>", ""]


def print_markdown(rep: Report) -> str:
    """Markdown for the HTML/PDF build: explicit heading ids, figures as <figure>, captions with anchors."""
    a = rep.front_anchors
    out = [f"# Document Control {{#{a[1]}}}", "", *doc_control(git("rev-parse", "--short", "HEAD")), "", f"# Table of Contents {{#{a[2]}}}", ""]
    out += nav([(h.level, h.text, h.anchor) for h in rep.headings if h.level <= 2], "contents")
    out += [f"# List of Figures {{#{a[3]}}}", "", *nav([(2, c.text, c.anchor) for c in rep.figures], "lof")]
    out += [f"# List of Tables {{#{a[4]}}}", "", *nav([(2, c.text, c.anchor) for c in rep.tables], "lot")]
    for pi, (_p, text) in enumerate(rep.parts):
        lines = text.split("\n")
        in_code, skip = False, set()
        for li, ln in enumerate(lines):
            if li in skip:
                continue
            if FENCE.match(ln):
                in_code = not in_code
                out.append(ln)
                continue
            if in_code:
                out.append(ln)
                continue
            anchor = rep.anchors.get((pi, li))
            if anchor:
                out.append(f"{ln} {{#{anchor}}}")
                continue
            im = IMG.match(ln.strip())
            if im:
                j = li + 1
                while j < len(lines) - 1 and not lines[j].strip() and j < li + 3:
                    j += 1
                cm = FIG_CAP.match(lines[j].strip()) if j < len(lines) else None
                fid, cap = "", ""
                if cm:
                    skip.add(j)
                    fid = f' id="fig-{cm.group(2).lower()}-{cm.group(3)}"'
                    cap = f"<figcaption>{inline_html(cm.group(1))}</figcaption>"
                kind = "shot" if "screenshots/" in im.group(2) else "diagram"
                out += ["", f'<figure class="{kind}"{fid}><img src="{html.escape(im.group(2))}" alt="{html.escape(im.group(1))}">{cap}</figure>', ""]
                continue
            cm = FIG_CAP.match(ln.strip())
            if cm:
                out += ["", f'<p class="caption fig-caption" id="fig-{cm.group(2).lower()}-{cm.group(3)}">{inline_html(cm.group(1))}</p>', ""]
                continue
            tm = TAB_CAP.match(ln.strip())
            if tm:
                out += ["", f'<p class="caption tab-caption" id="tab-{tm.group(2).lower()}-{tm.group(3)}">{inline_html(tm.group(1))}</p>', ""]
                continue
            out.append(ln)
        out.append("")
    return "\n".join(out)


CSS = """
@page { size: A4; margin: 20mm 17mm 18mm 17mm; }
html { font-family: 'Segoe UI', 'Helvetica Neue', Arial, sans-serif; font-size: 10.2pt; color: #0f172a; }
body { margin: 0; line-height: 1.48; }
h1 { font-size: 19pt; color: #1e1b4b; margin: 0 0 10pt; padding-bottom: 6pt; border-bottom: 2px solid #4f46e5;
     break-before: page; page-break-before: always; }
h2 { font-size: 13.5pt; color: #1e1b4b; margin: 15pt 0 5pt; break-after: avoid; page-break-after: avoid; }
h3 { font-size: 11.2pt; color: #312e81; margin: 11pt 0 4pt; break-after: avoid; page-break-after: avoid; }
h4, h5, h6 { font-size: 10.4pt; color: #312e81; margin: 9pt 0 3pt; break-after: avoid; }
p { margin: 0 0 6.5pt; text-align: justify; }
a { color: #3730a3; text-decoration: none; }
ul, ol { margin: 0 0 7pt 17pt; padding: 0; }
li { margin: 1.5pt 0; }
li > p { margin: 0; }
code { font-family: Consolas, 'Cascadia Mono', monospace; font-size: 8.5pt; background: #eef1f6; padding: 0 2.5px; border-radius: 3px;
       overflow-wrap: anywhere; }
pre { background: #f4f6fa; border: 1px solid #dde3eb; border-radius: 4px; padding: 6pt 8pt; font-size: 7.9pt; line-height: 1.32;
      white-space: pre-wrap; word-break: break-word; margin: 4pt 0 9pt; }
pre code { background: none; padding: 0; font-size: 7.9pt; }
table { border-collapse: collapse; width: 100%; margin: 3pt 0 11pt; font-size: 8.4pt; line-height: 1.32; }
thead { display: table-header-group; }
tr { break-inside: avoid; page-break-inside: avoid; }
th { background: #eef0ff; color: #1e1b4b; text-align: left; font-weight: 600; }
th, td { border: 1px solid #cbd5e1; padding: 2.6pt 4.5pt; vertical-align: top; overflow-wrap: break-word; }
td p, th p { text-align: left; margin: 0; }
tbody tr:nth-child(even) td { background: #f8fafc; }
figure { margin: 9pt 0 12pt; text-align: center; break-inside: avoid; page-break-inside: avoid; }
figure img { display: block; margin: 0 auto; max-width: 100%; max-height: 222mm; height: auto; }
figure.shot img { max-height: 150mm; border: 1px solid #cbd5e1; border-radius: 3px; }
figcaption, .fig-caption { font-size: 8.7pt; color: #334155; font-style: italic; margin-top: 4pt; text-align: center; }
.tab-caption { font-size: 8.9pt; font-weight: 600; color: #1e1b4b; margin: 10pt 0 2pt; break-after: avoid; page-break-after: avoid;
               text-align: left; }
blockquote { margin: 6pt 0 9pt; padding: 4pt 10pt; border-left: 3px solid #4f46e5; background: #f5f5ff; }
blockquote p { margin: 0; }
hr { border: none; border-top: 1px solid #cbd5e1; margin: 10pt 0; }
nav.toc a { display: flex; align-items: baseline; color: #0f172a; break-inside: avoid; }
nav.toc .toc-t { flex: 0 1 auto; }
nav.toc .toc-d { flex: 1 1 auto; border-bottom: 1px dotted #94a3b8; margin: 0 4pt; min-width: 12pt; transform: translateY(-2.5pt); }
nav.toc .toc-p { flex: 0 0 auto; min-width: 16pt; text-align: right; font-variant-numeric: tabular-nums; }
nav.toc a.toc-1 { font-weight: 600; margin-top: 5pt; font-size: 9.8pt; }
nav.toc a.toc-2 { font-size: 9pt; margin-left: 13pt; }
nav.lof a.toc-2, nav.lot a.toc-2 { margin-left: 0; margin-top: 1.5pt; }
.cover { height: 245mm; display: flex; flex-direction: column; justify-content: center; }
.cover-brand { font-size: 44pt; font-weight: 700; color: #1e1b4b; letter-spacing: -0.5pt; }
.cover-sub { font-size: 18pt; color: #4338ca; margin-top: 4pt; }
.cover-kind { font-size: 12.5pt; margin: 28pt 0 16pt; text-transform: uppercase; letter-spacing: 2.5pt; color: #475569; }
.cover-meta { width: 78%; font-size: 10.5pt; border: none; }
.cover-meta td { border: none; border-bottom: 1px solid #e2e8f0; padding: 5pt 4pt; background: none !important; }
.cover-meta td:first-child { color: #64748b; width: 38%; }
.cover-principle { margin-top: 34pt; font-size: 12.5pt; font-weight: 600; color: #15803d; }
.cover-note { margin-top: 10pt; font-size: 8.5pt; color: #64748b; }
@media screen { body { max-width: 190mm; margin: 0 auto; padding: 12mm 10mm; } .cover { height: auto; min-height: 240mm; } }
"""


def cover_html() -> str:
    rows = "".join(f"<tr><td>{html.escape(k)}</td><td>{html.escape(v)}</td></tr>" for k, v in META if k != "Project")
    return ('<div class="cover"><div class="cover-brand">SupportNova</div>'
            '<div class="cover-sub">Customer Complaint Resolution Intelligence</div>'
            '<div class="cover-kind">Project Report</div>'
            f'<table class="cover-meta">{rows}</table>'
            f'<div class="cover-principle">{PRINCIPLE}</div>'
            '<div class="cover-note">Lumora Home Technologies and all customers, orders, policies and complaints in this report are fictional.</div>'
            "</div>")


def to_html(md_text: str, with_cover: bool) -> str:
    import markdown  # installed in documentation/_build/.tools/pylib

    body = markdown.markdown(md_text, extensions=["extra", "sane_lists", "attr_list"], output_format="html5") if md_text else ""
    return (f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>{TITLE}</title><style>{CSS}</style></head>"
            f"<body>{cover_html() if with_cover else ''}{body}</body></html>")


# ----------------------------------------------------------------------------------------------- pdf
class Chrome:
    def __init__(self) -> None:
        self.profile = tempfile.mkdtemp(prefix="sn-pdf-")
        self.port = free_port()
        self.proc = subprocess.Popen([find_browser(), "--headless=new", f"--remote-debugging-port={self.port}",  # noqa: S603
                                      f"--user-data-dir={self.profile}", "--no-first-run", "--allow-file-access-from-files",
                                      "about:blank"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

    async def __aenter__(self) -> Chrome:
        for _ in range(150):
            with contextlib.suppress(Exception):
                page = next(t for t in httpx.get(f"http://127.0.0.1:{self.port}/json/list", timeout=2).json() if t["type"] == "page")
                break
            time.sleep(0.2)
        else:
            raise RuntimeError("browser did not start")
        self.ws = await websockets.connect(page["webSocketDebuggerUrl"], max_size=1024 * 1024 * 1024)
        self.cdp = CDP(self.ws)
        await self.cdp.send("Page.enable")
        return self

    async def __aexit__(self, *exc: object) -> None:
        with contextlib.suppress(Exception):
            await self.ws.close()
        self.proc.terminate()
        with contextlib.suppress(Exception):
            self.proc.wait(5)
        shutil.rmtree(self.profile, ignore_errors=True)

    async def pdf(self, html_file: Path, *, header: bool) -> bytes:
        url = html_file.as_uri()
        await self.cdp.send("Page.navigate", url=url)
        for _ in range(900):
            ok = await self.cdp.evaluate(f"location.href === {url!r} && document.readyState === 'complete' && "
                                         "[...document.images].every(i => i.complete)")
            if ok:
                break
            await asyncio.sleep(0.2)
        await self.cdp.evaluate("document.fonts ? document.fonts.ready.then(() => true) : true")
        params: dict[str, object] = {"printBackground": True, "preferCSSPageSize": True, "transferMode": "ReturnAsBase64"}
        if header:
            params.update(displayHeaderFooter=True,
                          headerTemplate=('<div style="font-size:7pt;color:#64748b;width:100%;margin:0 17mm;display:flex;'
                                          'justify-content:space-between;font-family:Segoe UI,Arial,sans-serif">'
                                          f'<span>{TITLE}</span><span>Version {VERSION} · {DATE:%d %B %Y}</span></div>'),
                          footerTemplate=('<div style="font-size:7pt;color:#64748b;width:100%;text-align:center;'
                                          'font-family:Segoe UI,Arial,sans-serif">Page <span class="pageNumber"></span> of '
                                          '<span class="totalPages"></span></div>'))
        res = await self.cdp.send("Page.printToPDF", **params)
        return base64.b64decode(res["data"])


def link_pages(pdf: bytes) -> dict[str, int]:
    import pymupdf

    pages: dict[str, int] = {}
    with pymupdf.open(stream=pdf, filetype="pdf") as d:
        for page in d:
            for ln in page.get_links():
                if ln.get("nameddest") and ln.get("page", -1) >= 0:
                    pages.setdefault(ln["nameddest"], ln["page"])
    return pages


def fill_pages(html_text: str, pages: dict[str, int], offset: int = 1) -> str:
    return re.sub(r'<span class="toc-p" data-a="([^"]+)">[^<]*</span>',
                  lambda m: f'<span class="toc-p" data-a="{m.group(1)}">'
                            f'{pages[m.group(1)] + offset if m.group(1) in pages else "000"}</span>', html_text)


async def build_pdf(rep: Report, body_md: str, out_pdf: Path, out_html: Path) -> int:
    import pymupdf

    work = Path(tempfile.mkdtemp(prefix="sn-report-"))
    body_file = DOC / "_build" / ".body.html"          # next to documentation/ so relative image paths resolve
    try:
        cover_file = work / "cover.html"
        cover_file.write_text(to_html("", with_cover=True), encoding="utf-8")
        body_html = to_html(body_md, with_cover=False).replace('src="diagrams/', 'src="../diagrams/').replace('src="../screenshots/', 'src="../../screenshots/')
        async with Chrome() as chrome:
            cover_pdf = await chrome.pdf(cover_file, header=False)
            pages: dict[str, int] = {}
            body_pdf = b""
            for _attempt in range(4):
                body_file.write_text(fill_pages(body_html, pages), encoding="utf-8")
                body_pdf = await chrome.pdf(body_file, header=True)
                new_pages = link_pages(body_pdf)
                if new_pages == pages:
                    break
                pages = new_pages
        missing = [h.anchor for h in rep.headings if h.level <= 2 and h.anchor not in pages]
        if missing:
            rep.issues.append(f"PDF: {len(missing)} TOC entries without a page number (first: {missing[:3]})")
        out = pymupdf.open()
        with pymupdf.open(stream=cover_pdf, filetype="pdf") as c:
            out.insert_pdf(c)
        with pymupdf.open(stream=body_pdf, filetype="pdf") as b:
            out.insert_pdf(b)
        with pymupdf.open(stream=body_pdf, filetype="pdf") as b:
            for t, a in zip(FRONT[1:], rep.front_anchors[1:], strict=True):
                hit = next((i for i, page in enumerate(b) if t in [x.strip() for x in page.get_text().splitlines()[:6]]), None)
                if a not in pages and hit is not None:
                    pages[a] = hit
        toc: list[list[object]] = [[1, "Cover", 1]]
        toc += [[1, t, pages[a] + 2] for t, a in zip(FRONT[1:], rep.front_anchors[1:], strict=True) if a in pages]
        for h in rep.headings:
            if h.level <= 2 and h.anchor in pages:
                level = min(h.level, int(toc[-1][0]) + 1)
                toc.append([level, re.sub(r"[`*]", "", h.text), pages[h.anchor] + 2])
        out.set_toc(toc)
        out.set_metadata({"title": TITLE, "author": "SupportNova project team",
                          "subject": "Customer Complaint Resolution Intelligence — project report",
                          "keywords": "SupportNova, ResponseX Intelligence, Generative AI PowerPlay, complaint resolution, validation",
                          "creator": "documentation/_build/build_report.py", "producer": "Chromium print + PyMuPDF"})
        out.save(out_pdf, garbage=3, deflate=True)
        n = out.page_count
        out.close()
        out_html.write_text(fill_pages(to_html(body_md, with_cover=True), pages), encoding="utf-8")
        return n
    finally:
        body_file.unlink(missing_ok=True)
        shutil.rmtree(work, ignore_errors=True)


# ----------------------------------------------------------------------------------------------- docx
def build_docx(rep: Report, out: Path) -> None:  # noqa: C901, PLR0912, PLR0915
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from PIL import Image

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.9)
    sec.different_first_page_header_footer = True
    usable_w, usable_h = Cm(17.2), Cm(21.5)
    st = doc.styles
    st["Normal"].font.name = "Calibri"
    st["Normal"].font.size = Pt(10.5)
    st["Normal"].paragraph_format.space_after = Pt(5)
    for name, size, color in (("Heading 1", 18, "1E1B4B"), ("Heading 2", 13.5, "1E1B4B"), ("Heading 3", 11.5, "312E81"),
                              ("Heading 4", 10.5, "312E81"), ("Heading 5", 10.5, "312E81"), ("Heading 6", 10.5, "312E81")):
        s = st[name]
        s.font.name, s.font.size = "Calibri", Pt(size)
        s.font.color.rgb = RGBColor.from_string(color)
    st["Heading 1"].paragraph_format.page_break_before = True
    for name in ("Figure Caption", "Table Caption"):
        s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = st["Caption"]
        s.font.size = Pt(9)
        s.font.color.rgb = RGBColor.from_string("334155" if name == "Figure Caption" else "1E1B4B")
        s.font.italic = name == "Figure Caption"
        s.font.bold = name == "Table Caption"
        s.paragraph_format.space_after = Pt(8 if name == "Figure Caption" else 3)
    st["Table Caption"].paragraph_format.keep_with_next = True
    code_style = st.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name, code_style.font.size = "Consolas", Pt(7.8)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.left_indent = Cm(0.2)

    def shade(el_pr, fill: str) -> None:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        el_pr.append(shd)

    def fld(run, kind: str) -> None:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        run._r.append(el)

    def instr(run, text: str) -> None:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = text
        run._r.append(el)

    def simple_field(par, code: str) -> None:
        r = par.add_run()
        fld(r, "begin")
        instr(r, code)
        fld(r, "separate")
        par.add_run("1")
        fld(par.add_run(), "end")

    def list_field(code: str, entries: list[tuple[int, str]], bold_level1: bool) -> None:
        """A TOC-type field pre-filled with its entries, so it is never empty; Word adds page numbers on update."""
        r = doc.add_paragraph().add_run()
        fld(r, "begin")
        instr(r, code)
        fld(r, "separate")
        for level, text in entries:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6 * (level - 1))
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(re.sub(r"[`*]", "", text))
            run.font.size = Pt(10 if level == 1 else 9.5)
            run.bold = level == 1 and bold_level1
        fld(doc.add_paragraph().add_run(), "end")

    hp = sec.header.paragraphs[0]
    hp.text = f"{TITLE} · Version {VERSION} · {DATE:%d %B %Y}"
    hp.runs[0].font.size = Pt(8)
    hp.runs[0].font.color.rgb = RGBColor.from_string("64748B")
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ")
    simple_field(fp, "PAGE")
    fp.add_run(" of ")
    simple_field(fp, "NUMPAGES")
    for r in fp.runs:
        r.font.size = Pt(8)
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)

    def inline(par, text: str, size: float | None = None, bold: bool = False, italic: bool = False) -> None:
        text = re.sub(r"<br\s*/?>", "\n", text)
        text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", lambda m: f"{m.group(1)} ({m.group(2)})" if m.group(2).startswith("http") else m.group(1), text)
        text = html.unescape(re.sub(r"</?(?:span|div|sup|sub|small|a|u)\b[^>]*>", "", text)).replace("\\|", "|")
        for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*\*(?![\w*]))", text):
            if not tok:
                continue
            b, it, mono = bold, italic, False
            if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
                tok, b = tok[2:-2], True
            elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
                tok, mono = tok[1:-1], True
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                tok, it = tok[1:-1], True
            run = par.add_run(tok)
            run.bold, run.italic = b, it
            if mono:
                run.font.name = "Consolas"
                run.font.size = Pt((size or 10.5) - 1.5)
                run.font.color.rgb = RGBColor.from_string("1E293B")
            elif size:
                run.font.size = Pt(size)

    def picture(src: str) -> None:
        p = (DOC / src).resolve()
        if p.suffix.lower() == ".svg":
            p = p.with_suffix(".png")
        if not p.exists():
            doc.add_paragraph(f"[missing image: {src}]")
            return
        with Image.open(p) as im:
            w, h = im.size
        limit_h = Cm(14) if "screenshots" in src else usable_h
        width = usable_w
        if width * h / w > limit_h:
            width = int(limit_h * w / h)
        doc.add_picture(str(p), width=width)
        par = doc.paragraphs[-1]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True

    def table(rows: list[list[str]], meta: bool = False) -> None:
        ncols = max(len(r) for r in rows)
        t = doc.add_table(rows=0, cols=ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        size = 7.8 if ncols >= 6 else 8.4 if ncols >= 4 else 9.2
        for i, row in enumerate(rows):
            tr = t.add_row()
            header = i == 0 and not meta
            if header:
                hdr = OxmlElement("w:tblHeader")
                hdr.set(qn("w:val"), "true")
                tr._tr.get_or_add_trPr().append(hdr)
            for jx in range(ncols):
                par = tr.cells[jx].paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                inline(par, row[jx] if jx < len(row) else "", size=size, bold=header)
                if header:
                    shade(tr.cells[jx]._tc.get_or_add_tcPr(), "EEF0FF")
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---- cover (first page: no header / footer)
    for text, size, color, bold, space in (("SupportNova", 38, "1E1B4B", True, 4), ("Customer Complaint Resolution Intelligence", 17, "4338CA", False, 26),
                                           ("PROJECT REPORT", 12, "475569", False, 14)):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space)
        if text == "SupportNova":
            p.paragraph_format.space_before = Pt(110)
        run = p.add_run(text)
        run.font.size, run.bold = Pt(size), bold
        run.font.color.rgb = RGBColor.from_string(color)
    table([[k, v] for k, v in META if k != "Project"], meta=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    run = p.add_run(PRINCIPLE)
    run.bold, run.font.size = True, Pt(12.5)
    run.font.color.rgb = RGBColor.from_string("15803D")
    note = doc.add_paragraph().add_run("Lumora Home Technologies and all customers, orders, policies and complaints in this report are fictional.")
    note.font.size, note.font.color.rgb = Pt(8.5), RGBColor.from_string("64748B")

    # ---- front matter
    doc.add_heading("Document Control", level=1)
    dc = [[c.strip() for c in ln.strip("|").split(" | ")] for ln in doc_control(git("rev-parse", "--short", "HEAD"))]
    table([dc[0], *dc[2:]])
    doc.add_heading("Table of Contents", level=1)
    list_field('TOC \\o "1-2" \\h \\z \\u', [(h.level, h.text) for h in rep.headings if h.level <= 2], True)
    doc.add_heading("List of Figures", level=1)
    list_field('TOC \\h \\z \\t "Figure Caption,1"', [(1, c.text) for c in rep.figures], False)
    doc.add_heading("List of Tables", level=1)
    list_field('TOC \\h \\z \\t "Table Caption,1"', [(1, c.text) for c in rep.tables], False)

    # ---- body
    block_start = re.compile(r"^(#{1,6} |\||```|~~~|[-*+] |\d+\. |!\[|>|\*Figure |\*\*Table )")
    for _p, text in rep.parts:
        lines = text.split("\n")
        i = 0
        while i < len(lines):
            ln = lines[i]
            s = ln.strip()
            if FENCE.match(ln):
                i += 1
                code = []
                while i < len(lines) and not FENCE.match(lines[i]):
                    code.append(lines[i])
                    i += 1
                i += 1
                par = doc.add_paragraph(style="Code Block")
                par.add_run("\n".join(code))
                shade(par._p.get_or_add_pPr(), "F4F6FA")
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
                continue
            h = HEAD.match(s)
            if h:
                doc.add_heading(re.sub(r"[`*]", "", h.group(2)), level=min(len(h.group(1)), 6))
                i += 1
                continue
            if not s or re.fullmatch(r"-{3,}|\*{3,}", s):
                i += 1
                continue
            im = IMG.match(s)
            if im:
                picture(im.group(2))
                i += 1
                continue
            cm = FIG_CAP.match(s)
            if cm:
                par = doc.add_paragraph(style="Figure Caption")
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                inline(par, cm.group(1))
                i += 1
                continue
            tm = TAB_CAP.match(s)
            if tm:
                inline(doc.add_paragraph(style="Table Caption"), tm.group(1))
                i += 1
                continue
            if s.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    cells = [c.strip() for c in re.split(r"(?<!\\)\|", lines[i].strip().strip("|"))]
                    if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                        rows.append(cells)
                    i += 1
                if rows:
                    table(rows)
                continue
            lm = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", ln)
            if lm:
                depth = min(len(lm.group(1).replace("\t", "    ")) // 2, 3)
                bullet = lm.group(2) in "-*+"
                par = doc.add_paragraph(style="List Bullet" if bullet else "List Paragraph")
                par.paragraph_format.left_indent = Cm(0.63 + 0.6 * depth)
                par.paragraph_format.space_after = Pt(2)
                item = lm.group(3)
                i += 1
                while i < len(lines) and lines[i].strip() and lines[i].startswith("  ") and not block_start.match(lines[i].strip()):
                    item += " " + lines[i].strip()
                    i += 1
                if not bullet:
                    par.paragraph_format.first_line_indent = Cm(-0.5)
                    item = f"{lm.group(2)}  {item}"
                inline(par, item)
                continue
            if s.startswith(">"):
                buf = []
                while i < len(lines) and lines[i].strip().startswith(">"):
                    buf.append(lines[i].strip().lstrip(">").strip())
                    i += 1
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.6)
                shade(par._p.get_or_add_pPr(), "F5F5FF")
                inline(par, " ".join(buf), italic=True)
                continue
            if s.startswith("<"):
                i += 1
                continue
            buf = [s]
            i += 1
            while i < len(lines) and lines[i].strip() and not block_start.match(lines[i].strip()):
                buf.append(lines[i].strip())
                i += 1
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            inline(par, " ".join(buf))
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Customer Complaint Resolution Intelligence — project report"
    doc.core_properties.author = "SupportNova project team"
    doc.save(out)


# ----------------------------------------------------------------------------------------------- main
def main() -> None:
    rep = load()
    words = len(re.findall(r"\b\w+\b", re.sub(r"```.*?```", "", "\n".join(t for _p, t in rep.parts), flags=re.S)))
    summary = (f"{len(rep.parts)} parts, {len(rep.headings)} headings, {len(rep.figures)} figures, {len(rep.tables)} tables, "
               f"{words} words outside code blocks")
    if "--check" not in sys.argv:
        (DOC / "SupportNova_Project_Report.md").write_text(public_markdown(rep), encoding="utf-8")
        if "--no-pdf" not in sys.argv:
            n = asyncio.run(build_pdf(rep, print_markdown(rep), DOC / "SupportNova_Project_Report.pdf", DOC / "SupportNova_Project_Report.html"))
            summary += f", PDF {n} pages"
        if "--no-docx" not in sys.argv:
            build_docx(rep, DOC / "SupportNova_Project_Report.docx")
            summary += ", DOCX written"
    (DOC / "_build" / "report-check.txt").write_text("\n".join([summary, *rep.issues]) + "\n", encoding="utf-8")
    print("\n".join(rep.issues[:80]))
    if len(rep.issues) > 80:
        print(f"... {len(rep.issues) - 80} more in _build/report-check.txt")
    print(summary)


if __name__ == "__main__":
    main()
