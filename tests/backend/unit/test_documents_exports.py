"""Document ingestion (PDF/DOCX/MD parsing, sections, chunks, facts, metadata) and report exports."""

from __future__ import annotations

import csv
import io
from pathlib import Path

import pytest
import yaml

from supportnova.core.errors import ValidationFailed
from supportnova.document_processing.chunking import chunk_sections
from supportnova.document_processing.facts import diff_versions, extract_facts, facts_to_json
from supportnova.document_processing.parsers import ParsedSection, parse_document
from supportnova.document_processing.validation import validate_metadata, version_key
from supportnova.reporting.exports import Report, Section, Table, csv_safe, render, tabular
from supportnova.security import injection

ROOT = Path(__file__).resolve().parents[3]
KB = ROOT / "knowledge_base"
MANIFEST = yaml.safe_load((KB / "manifest.yaml").read_text(encoding="utf-8"))


def _file(entry: dict) -> Path:  # type: ignore[type-arg]
    return KB / entry["file"]


def test_knowledge_base_meets_srs_minimums() -> None:
    docs = {e["doc_id"] for e in MANIFEST}
    formats = {e["format"] for e in MANIFEST}
    assert len(docs) >= 20
    assert {"pdf", "docx"} <= formats
    statuses = {e["status"] for e in MANIFEST}
    assert "Active" in statuses and ({"Previous", "Superseded"} & statuses)


@pytest.mark.parametrize("fmt", ["pdf", "docx"])
def test_parse_sections_from_real_documents(fmt: str) -> None:
    entry = next(e for e in MANIFEST if e["format"] == fmt and e["status"] == "Active")
    parsed = parse_document(_file(entry).read_bytes(), fmt)
    assert parsed.sections, entry["doc_id"]
    assert all(s.section_id and s.heading for s in parsed.sections)
    assert parsed.detected_metadata.get("doc_id") == entry["doc_id"]
    chunks = chunk_sections(entry["doc_id"], str(entry["version"]), parsed.sections)
    assert chunks and all(c.chunk_uid.startswith(f"{entry['doc_id']}@") for c in chunks)


def test_pdf_keeps_repeated_body_text_and_reads_tables_by_row() -> None:
    """Only running headers/footers are dropped; a value repeated in the body (a department name in a table)
    is content. Tables stay one row per line - the routing table used to lose its department cells."""
    from fpdf import FPDF

    class Doc(FPDF):
        def header(self) -> None:
            self.set_font("helvetica", size=8)
            self.cell(0, 6, "TST-POL-01 v1.0 - Test Policy", align="C")
            self.ln(20)

        def footer(self) -> None:
            self.set_y(-14)
            self.set_font("helvetica", size=8)
            self.cell(0, 6, f"Page {self.page_no()} of {{nb}}", align="C")

    pdf = Doc()
    for n in range(1, 4):
        pdf.add_page()
        pdf.set_font("helvetica", "B", 13)
        pdf.cell(0, 8, f"{n} Section {n}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("helvetica", size=10)
        pdf.multi_cell(0, 6, "Billing Operations (DEPT-BIL) handles this.", new_x="LMARGIN", new_y="NEXT")
        if n == 2:
            with pdf.table() as table:
                for cells in (("Rule ID", "Subcategory", "Primary department"),
                              ("RTE-004", "BIL-DUP Duplicate Charge", "Billing Operations (DEPT-BIL)")):
                    row = table.row()
                    for cell in cells:
                        row.cell(cell)
    parsed = parse_document(bytes(pdf.output()), "pdf")
    assert "TST-POL-01 v1.0 - Test Policy" not in parsed.full_text and "Page 2 of 3" not in parsed.full_text
    assert parsed.full_text.count("Billing Operations (DEPT-BIL) handles this.") == 3
    assert "RTE-004 | BIL-DUP Duplicate Charge | Billing Operations (DEPT-BIL)" in parsed.full_text
    assert [s.section_id for s in parsed.sections] == ["1", "2", "3"]


def test_facts_and_version_diff() -> None:
    old = [ParsedSection("3.1", "Standard Refund Window", 3, "Products may be returned within 30 calendar days of delivery.")]
    new = [ParsedSection("3.1", "Standard Refund Window", 3, "Products may be returned within 21 calendar days of delivery.")]
    as_dict = lambda secs: [{"section_id": x.section_id, "heading": x.heading, "text": x.text} for x in secs]  # noqa: E731
    changes = diff_versions(as_dict(old), as_dict(new), facts_to_json(extract_facts(old)), facts_to_json(extract_facts(new)))
    assert changes and changes[0]["section_id"] == "3.1"
    assert any(float(v["old"]) == 30 and float(v["new"]) == 21 for v in changes[0].get("value_changes", []))


def test_malicious_document_sections_flagged() -> None:
    path = KB / "security_samples" / "MAL-DOC-99_v1.0.docx"
    parsed = parse_document(path.read_bytes(), "docx")
    flagged = [s.section_id for s in parsed.sections if injection.scan(s.text).is_suspicious]
    assert flagged, "embedded instructions in the malicious sample must be detected"
    clean = [s for s in parsed.sections if s.section_id not in flagged]
    assert clean, "legitimate sections of the same document stay usable"


def test_metadata_validation_and_version_order() -> None:
    with pytest.raises(ValidationFailed) as err:
        validate_metadata({"doc_id": "BAD ID", "title": "", "doc_type": "policy", "version": "x", "status": "Weird"})
    fields = {d["field"] for d in err.value.details}  # type: ignore[union-attr]
    assert {"version", "status"} <= fields
    ok = validate_metadata({"doc_id": "REF-POL-02", "title": "Refund Policy", "doc_type": "policy", "version": "2.1",
                            "status": "Active", "effective_date": "2026-09-15"})
    assert ok.doc_id == "REF-POL-02"
    assert version_key("2.10") > version_key("2.9") > version_key("2.0")


def test_corrupted_document_is_rejected_cleanly() -> None:
    from supportnova.core.errors import AppError
    with pytest.raises(AppError):
        parse_document(b"%PDF-1.4 this is not really a pdf", "pdf")


# ---------------------------------------------------------------- exports
def test_csv_neutralises_formula_injection() -> None:
    content, media, ext = tabular("t", ["a", "b"], [["=HYPERLINK(\"http://x\")", "+1"], ["safe", "-2"]], "csv")
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert rows[1][0].startswith("'=") and rows[1][1].startswith("'+") and rows[2][1] == "'-2"
    assert ext == "csv" and media.startswith("text/csv")
    assert csv_safe("@cmd") == "'@cmd"


def test_xlsx_never_contains_formulas() -> None:
    from openpyxl import load_workbook
    content, _media, ext = tabular("t", ["a"], [["=1+1"], ["plain"]], "xlsx")
    wb = load_workbook(io.BytesIO(content))
    ws = wb[wb.sheetnames[0]]
    assert ext == "xlsx"
    assert ws["A2"].data_type == "s" and ws["A2"].value == "=1+1"


def test_pdf_report_renders_unicode_safely() -> None:
    report = Report("Déjà vu — report", [Section("Summary", metrics=[("Complaints", 3)], paragraphs=["Quotes “smart” and ≥ signs"],
                                                  tables=[Table(["x", "y"], [["ü", 1.5]])])], banner="Important notice")
    content, media, ext = render(report, "pdf")
    assert content.startswith(b"%PDF") and ext == "pdf" and media == "application/pdf"


def test_json_render_roundtrip() -> None:
    content, _m, ext = render(Report("r", [Section("s", metrics=[("k", 1)])]), "json")
    assert ext == "json" and b'"metrics"' in content
